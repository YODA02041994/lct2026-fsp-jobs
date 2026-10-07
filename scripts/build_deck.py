#!/usr/bin/env python3
"""Собирает презентацию «Карьера ФСП» из шаблона организатора (docs/materials/ЛЦТ_2026_Шаблон_презентации_ФСП.pptx).

Обязательные слайды 7–11 заполняются по месту (макеты не меняются), для содержания берутся образцы макетов 12–29 шаблона
(заполняем их заполнители и диаграммы), остальные слайды-инструкции удаляются. Снимки экрана — из docs/screens/, если есть.

    python3 scripts/build_deck.py            # → docs/build/Карьера_ФСП_презентация.pptx (+ копия в frontend/files/)
"""
from __future__ import annotations
import copy
import os
import shutil

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.util import Pt, Inches

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "docs", "materials", "ЛЦТ_2026_Шаблон_презентации_ФСП.pptx")
OUT = os.path.join(ROOT, "docs", "build", "Карьера_ФСП_презентация.pptx")
SCREENS = os.path.join(ROOT, "docs", "screens")
STAND = "docsemenov.ru/fsp"


# ---------------------------------------------------------------- помощники
def ph(slide, idx):
    """Заполнитель по idx; если образец слайда его не содержит — клонируем из макета."""
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == idx:
            return shape
    for lp in slide.slide_layout.placeholders:
        if lp.placeholder_format.idx == idx:
            slide.shapes.clone_placeholder(lp)
            return ph(slide, idx)
    raise KeyError(f"нет заполнителя {idx} ни на слайде, ни в макете")


def by_id(slide, shape_id):
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape
    raise KeyError(f"нет формы {shape_id}")


def remove(shape):
    el = shape._element
    el.getparent().remove(el)


def set_text(shape, lines, size=None, bold=None, color=None, align=None, bullets=False):
    """Пишет строки в текстовую рамку, перенося свойства первого прогона шаблона целиком (rPr: шрифт, цвет темы, размер).
    Так текст наследует цвет шаблона — тёмный на белых карточках, белый на фиолетовом."""
    if isinstance(lines, str):
        lines = [lines]
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    r0 = p0.runs[0] if p0.runs else None
    rPr = copy.deepcopy(r0._r.rPr) if r0 is not None and r0._r.rPr is not None else None
    pPr = copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    base_align = p0.alignment
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if pPr is not None and i > 0:
            p._p.insert(0, copy.deepcopy(pPr))
        if align is not None:
            p.alignment = align
        elif base_align is not None:
            p.alignment = base_align
        run = p.add_run()
        if rPr is not None:
            run._r.insert(0, copy.deepcopy(rPr))
        run.text = ("• " if bullets else "") + line
        f = run.font
        if size is not None:
            f.size = Pt(size)
        if bold is not None:
            f.bold = bold
        if color is not None:
            f.color.rgb = color if isinstance(color, RGBColor) else RGBColor.from_string(color)


def set_title(slide, text, color=None):
    """Заголовок слайда; розовая плашка шаблона (скруглённый прямоугольник слева сверху) растягивается под текст."""
    set_text(ph(slide, 0), text, color=color)
    for shape in slide.shapes:
        if shape.is_placeholder or shape.shape_type == 13:
            continue
        if shape.top < Inches(0.6) and shape.left < Inches(0.6) and shape.height < Inches(1.0) and shape.width > Inches(2):
            need = Inches(0.6 + 0.19 * len(text))   # заглавные Montserrat жирным ≈ 0,19 дюйма на знак при кегле шаблона
            if shape.width < need:
                shape.width = need


def first_existing(*names):
    for n in names:
        path = os.path.join(SCREENS, n)
        if os.path.exists(path):
            return path
    return None


def picture(slide, idx, *filenames):
    """Вставляет первый существующий снимок в заполнитель картинки; если снимков нет — заполнитель убираем."""
    path = first_existing(*filenames)
    shape = ph(slide, idx)
    if path:
        shape.insert_picture(path)
        return True
    remove(shape)
    return False


def keep_slides(prs, order_1based):
    """Оставляет только перечисленные слайды в заданном порядке."""
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    keep = [ids[i - 1] for i in order_1based]
    for sld in ids:
        if sld not in keep:
            prs.part.drop_rel(sld.rId)
            sldIdLst.remove(sld)
    for sld in keep:
        sldIdLst.remove(sld)
    for sld in keep:
        sldIdLst.append(sld)


def chart_of(slide):
    for shape in slide.shapes:
        if shape.has_chart:
            return shape.chart
    raise KeyError("на слайде нет диаграммы")


VIOLET = "310F53"
PINK = "FF0053"
DARK = "1C1D22"


# ---------------------------------------------------------------- содержание
def main():
    prs = Presentation(TEMPLATE)
    S = {i: prs.slides[i - 1] for i in range(1, len(prs.slides) + 1)}

    # 7 — титул (обязательный): название команды + название задачи (уже в шаблоне), логотипы постановщика оставлены
    set_text(ph(S[7], 0), "ПИРОГОВСКИЙ УНИВЕРСИТЕТ")
    set_text(ph(S[7], 12), "Цифровая платформа-агрегатор ИТ-вакансий с верифицированным профилем достижений участника "
                           "Федерации спортивного программирования (ФСП) · решение «Карьера ФСП»")

    # 8 — о команде (обязательный)
    s = S[8]
    set_title(s, "КАРЬЕРА ФСП")
    set_text(by_id(s, 5), ["Платформа найма наоборот: работодатель описывает потребность и получает подборку кандидатов с "
                           "грейдом, подтверждённым тестом, — и сам направляет приглашение с зарплатой. Контакты открываются "
                           "после согласия кандидата. Профиль связан с ФСП ID и подтверждёнными достижениями."], size=11)
    set_text(by_id(s, 8), ["Грейд не по резюме, а по тесту из заданий с параметрами — утечка ответов бессмысленна. "
                           "Формула подбора найдена эволюционным поиском и проверена на мирах, которых не видела. "
                           "Каждая оценка — объяснимая вероятность «подходит» с обоснованием."], size=11)
    set_text(by_id(s, 14), ["Капитан: Александр Гаданов",
                            "Участников: 3 человека",
                            "Команда ЛЦТ-2026: та же тройка решала задачу № 4 основного трека (контроль качества денситометрии)",
                            "Место работы и учёбы: РНИМУ им. Н. И. Пирогова и практическая разработка",
                            "Город и регион: Москва"], size=11)
    try:
        remove(ph(s, 10))      # пустой заполнитель фотографии команды
    except KeyError:
        pass

    # 9 — участники (обязательный): три карточки, лишние удаляем; телефоны не публикуем
    s = S[9]
    set_title(s, "КОМАНДА")
    cards = [
        (15, 9, "Александр Гаданов", ["Капитан, лидер команды на платформе", "продуктовые сценарии, сдача, презентация", "ник в мессенджере: по запросу", "место работы/учёбы: (заполнить)"]),
        (58, 57, "Андрей Семенов", ["Архитектура, бэкенд, механики подбора и тестирования, валидация", "Telegram: @drsav1994", "РНИМУ им. Н. И. Пирогова, к.м.н.", "детский травматолог-ортопед, разработчик"]),
        (61, 60, "Саак Серобян", ["Тестирование прототипа, пилот теста, интерфейс", "ник в мессенджере: по запросу", "место работы/учёбы: (заполнить)"]),
    ]
    for name_id, role_id, name, lines in cards:
        set_text(by_id(s, name_id), name)
        set_text(by_id(s, role_id), lines, size=9)
    for sid in (62, 63, 64, 65, 66, 67, 5, 6):
        try:
            remove(by_id(s, sid))
        except KeyError:
            pass

    # 10 — история команды (обязательный)
    s = S[10]
    set_title(s, "ИСТОРИЯ КОМАНДЫ")
    set_text(by_id(s, 37), ["Собрались на ЛЦТ-2026 ради задачи № 4 (ИИ-контроль качества денситометрии) и сдали её 29 сентября; "
                            "спецтрек ФСП взяли второй задачей той же командой: интересно было проверить, можно ли вместо "
                            "«резюме и собеседование» сделать «тест и подтверждённый профиль»."], size=11)
    set_text(by_id(s, 43), ["Задача про честность данных: резюме не подтверждает компетентность, а у участников ФСП есть то, "
                            "чего нет ни у кого, — проверяемая история результатов. Это и есть продукт, который хочется построить."], size=11)
    set_text(by_id(s, 40), ["Главная сложность — проверить механики без живых пользователей за шесть дней. Построили синтетические миры "
                            "с известной истиной, считали метрики из ТЗ и записывали результаты как есть, включая проигрыш первой "
                            "формулы «словам резюме» в честном мире. Затем формулу нашли эволюционным поиском и перепроверили на новых мирах."], size=11)

    # 11 — коротко о решении (обязательный)
    s = S[11]
    set_text(ph(s, 38), ["FastAPI + PostgreSQL + Vue 3 без сборки; Docker Compose; OpenAPI на 55 операций; 644 автотеста.",
                         "Тест: 136 шаблонов заданий с параметрами, модель ответа с угадыванием, адаптивные блоки, EAP; грейд из θ.",
                         "Подбор: шесть объяснимых слагаемых, зарплатный множитель и покрытие стека; оценка — калиброванная вероятность «подходит».",
                         "Формула найдена эволюцией (модель — оператор мутации, программа — фитнес) и упрощена; в пределах шума от бустинга.",
                         "ФСП ID по OpenID Connect (Keycloak), реестр с уровнями доверия, события из календаря Минспорта."], size=10.5)
    set_text(ph(s, 51), "Рынок и развитие")
    set_text(ph(s, 42), ["Для работодателя: подборка по подтверждённому грейду вместо потока откликов; обоснование у каждого кандидата; "
                         "контакт только по согласию — меньше спама обеим сторонам.",
                         "Для участника ФСП: достижения работают на карьеру — подтверждённая запись в профиле видна работодателю.",
                         "Дальше: пилот с живыми кандидатами, подключение реестра ФСП и Keycloak, калибровка подбора на реальных "
                         "откликах и приглашениях, интеграция с ATS, адаптивная версия."], size=10.5)

    # 24 — задача и механика (проблема → решение)
    s = S[24]
    set_title(s, "НАЙМ НАОБОРОТ")
    for idx, t in ((33, "Проблема"), (34, "Что сделали"), (35, "Что это даёт")):
        set_text(ph(s, idx), t, size=14, bold=True, color=VIOLET)
    for idx in (26, 31, 32):
        body = ph(s, idx); body.top = Inches(4.25); body.height = Inches(2.7)
    set_text(ph(s, 26), ["Резюме не подтверждает компетентность, отклики идут потоком, грейд в вакансии и грейд кандидата — "
                         "разные вещи. У участников ФСП есть проверяемые результаты, но они не работают на найм."], size=11, color=DARK)
    set_text(ph(s, 31), ["Кандидат проходит опрос и тест и получает категорию «специализация × грейд». Работодатель описывает "
                         "потребность, получает подборку с обоснованием и сам направляет приглашение с зарплатой."], size=11, color=DARK)
    set_text(ph(s, 32), ["Контакты открываются только после согласия кандидата. Подборки хранятся версиями. Достижения ФСП — "
                         "подтверждённое слагаемое оценки, а не строка в резюме."], size=11, color=DARK)

    # 25 — путь пользователей (стадии)
    s = S[25]
    set_title(s, "ПУТЬ ПОЛЬЗОВАТЕЛЕЙ")
    steps = [("Регистрация и опрос", "кандидат: почта, специализация, стек, зарплата, формат; резюме PDF разбирается в поля"),
             ("Тест на грейд", "10 заданий с параметрами, два адаптивных блока, около 10 минут; подтверждённые технологии"),
             ("Категория и профиль", "грейд из θ, ФСП ID и достижения с уровнем доверия; приватность по полям"),
             ("Потребность → подборка", "работодатель описывает задачу; подборка по категориям с оценкой и обоснованием"),
             ("Приглашение → контакт", "зарплата «от–до» в приглашении; согласие кандидата открывает контакты")]
    for (t, d), (ti, di) in zip(steps, [(26, 27), (28, 29), (30, 31), (32, 33), (34, 35)]):
        set_text(ph(s, ti), t, size=12, bold=True, color=VIOLET)
        set_text(ph(s, di), d, size=9.5, color=DARK)

    # 15 — механика подбора (пункты)
    s = S[15]
    set_title(s, "ПОДБОР: ШЕСТЬ ОБЪЯСНИМЫХ СЛАГАЕМЫХ", color=VIOLET)
    rows = ["Стек (0,50): подтверждённая тестом технология — 0,85–1,0; заявленная — не ниже 0,55; проваленная по 1–2 заданиям — 0,3, а не ноль",
            "Грейд (0,22): подтверждённый тестом; соседний — 0,6; заявленный кандидатом — только подсказка, если соседний",
            "Зарплата — множитель: ожидания выше вилки работодателя опускают кандидата — самый сильный фактор формулы",
            "Покрытие стека: ≥ 90 % требуемого — ×1,08; ≤ 40 % — ×0,78; сила профиля, ФСП, активность — добавка",
            "Итог — калиброванная вероятность «подходит»; обоснование одной строкой; работодатель может менять веса"]
    for idx, text in zip((15, 16, 17, 18, 19), rows):
        set_text(ph(s, idx), text, size=11.5, color=DARK)

    # 21 — подбор против слов резюме (столбцы) + три числа
    s = S[21]
    set_title(s, "ПОДБОР ПРОТИВ СЛОВ РЕЗЮМЕ")
    cd = CategoryChartData(); cd.categories = ["честные резюме", "типичные", "приукрашенные"]
    cd.add_series("подбор, P@5", (0.97, 0.88, 0.77)); cd.add_series("слова резюме, P@5", (0.80, 0.83, 0.65))
    chart_of(s).replace_data(cd)
    for a, b, big, small in ((21, 18, "0,97", "доля подходящих среди первых пяти в честном мире (слова резюме — 0,80); тест в мирах — настоящий движок"),
                             (22, 23, "+0,14…+0,16", "выигрыш по nDCG@10 во всех трёх мирах честности резюме"),
                             (24, 25, "8 из 8", "вариантов генератора, где преимущество над словами резюме значимо: от +0,10 до +0,23")):
        set_text(ph(s, a), big, bold=True, color=PINK); set_text(ph(s, b), small, size=9.5, color=DARK)

    # 22 — формула найдена эволюцией (полосы) + четыре числа
    s = S[22]
    set_title(s, "ФОРМУЛА НАЙДЕНА ЭВОЛЮЦИЕЙ")
    cd = CategoryChartData(); cd.categories = ["исходная", "слова", "заход 3", "итоговая", "потолок"]
    cd.add_series("фитнес на контрольных мирах, реалистичный тест", (0.591, 0.682, 0.679, 0.804, 0.813))
    chart_of(s).replace_data(cd)
    for a, b, big, small in ((21, 18, "221", "обращение к DeepSeek V4.1 Flash за пять заходов — модель предлагает изменения, оценивает программа-фитнес"),
                             (22, 23, "$0,23", "вся эволюция; порог шума отсеял 52 из 60 потомков последнего захода, упрощение убрало 18 правил"),
                             (24, 25, "0,591 → 0,804", "на мирах, не участвовавших в поиске; потолок бустинга 0,813 — разница в пределах шума"),
                             (26, 27, "честно", "формула дня проиграла словам резюме, когда тест в мирах стал реальным, — это и нашла проверка")):
        set_text(ph(s, a), big, bold=True, color=PINK); set_text(ph(s, b), small, size=9.5, color=DARK)

    # 20 — тест на грейд (пункты)
    s = S[20]
    set_title(s, "ТЕСТ НА ГРЕЙД", color=VIOLET)
    rows = ["136 шаблонов для восьми специализаций; у каждого кандидата свой экземпляр — числа, имена, данные другие",
            "Модель ответа с угадыванием, популяционный априор, EAP, два адаптивных блока по пять заданий, около десяти минут",
            "Утечка 40 экземпляров не меняет оценку (θ̂ −1,54 → −1,54); фиксированный набор пропускает слабых втрое чаще",
            "Технология подтверждена, если по ней ≥ 50 % верных; подтверждённый грейд меняется не чаще раза в 90 дней",
            "Банк: 96 рукописных шаблонов и 40 от модели — через валидатор и четыре ревью по документации"]
    for idx, text in zip((15, 16, 17, 18, 19), rows):
        set_text(ph(s, idx), text, size=11.5, color=DARK)
    shot = first_existing("16_test_question.png", "candidate_test.png")
    if shot:
        box = [sh for sh in s.shapes if not sh.is_placeholder and sh.left < Inches(1) and sh.top > Inches(1) and sh.width > Inches(5)]
        if box:
            b = box[0]
            s.shapes.add_picture(shot, b.left + Inches(0.15), b.top + Inches(0.15), width=b.width - Inches(0.3))

    # 23 — валидация теста (кольцо) + четыре числа
    s = S[23]
    set_title(s, "ВАЛИДАЦИЯ ТЕСТА НА СИНТЕТИКЕ")
    cd = CategoryChartData(); cd.categories = ["точный грейд — 58 %", "соседний — 40 %", "ошибка на два — 2 %"]
    cd.add_series("доля кандидатов", (0.58, 0.40, 0.02))
    ch = chart_of(s); ch.replace_data(cd)
    from pptx.enum.chart import XL_LEGEND_POSITION
    ch.plots[0].has_data_labels = False          # подписи внутри кольца обрезались — числа перенесены в легенду
    from pptx.oxml.ns import qn
    for ser in ch.plots[0].series:               # подписи заданы и на уровне ряда — убираем и там
        for dl in ser._element.findall(qn("c:dLbls")):
            ser._element.remove(dl)
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(12); ch.legend.font.color.rgb = RGBColor.from_string(DARK)
    for a, b, big, small in ((21, 18, "0,98", "грейд в пределах соседнего (10 заданий; 2 зерна × 300 синтетических кандидатов восьми специализаций)"),
                             (22, 23, "0,58 / 0,62", "точный грейд при 10 / 15 заданиях; тест ограничен объёмом информации, а не настройкой"),
                             (24, 25, "0,79–0,82", "AUC между соседними грейдами junior/middle и middle/senior"),
                             (26, 27, "0,50 → 0,59", "согласованность повтора при 10 → 15 заданиях; путь улучшения — длиннее тест и больше шаблонов")):
        set_text(ph(s, a), big, bold=True, color=PINK); set_text(ph(s, b), small, size=9.5, color=DARK)

    # 16 — профиль ФСП (сетка 6)
    s = S[16]
    set_title(s, "ПРОФИЛЬ ДОСТИЖЕНИЙ ФСП")
    grid = [("ФСП ID", "вход по OpenID Connect, совместимо с Keycloak; в прототипе — издатель-заглушка с тем же протоколом; связь можно разорвать"),
            ("Реестр и доверие", "146 событий из календаря Минспорта; результаты → достижения A1–A5 со сроком действия; уровни «источник / документ / заявлено» не смешиваются"),
            ("Нормы и сроки", "звания бессрочно, КМС три года, разряды два; разряд по месту и уровню соревнования (приказы Минспорта)"),
            ("В подборе", "кандидат без истории ФСП проходит тот же путь — слагаемое достижений ноль; достижение решает ничью, а не перебивает стек")]
    for (t, d), (ni, ti, di) in zip(grid, [(49, 37, 38), (50, 39, 40), (51, 41, 42), (52, 43, 44)]):
        set_text(ph(s, ni), f"0{grid.index((t, d)) + 1}", bold=True, color=PINK); set_text(ph(s, ti), t, bold=True, color=VIOLET); set_text(ph(s, di), d, size=9.5, color=DARK)

    # 12 — архитектура и техника
    s = S[12]
    set_title(s, "АРХИТЕКТУРА И ТЕХНИКА")
    set_text(ph(s, 1), ["Фронтенд: Vue 3 без сборки, два кабинета, работает без интернета",
                        "Бэкенд: FastAPI, SQLAlchemy, Pydantic; OpenAPI на 55 операций с кодами ошибок",
                        "Хранение: PostgreSQL 16 (стенд, Docker), SQLite для локального запуска",
                        "Docker Compose: db / api / web; запуск одной командой; 644 автотеста",
                        "152-ФЗ: три согласия, роли, Argon2, JWT, журнал действий, приватность по полям",
                        "Модель DeepSeek — только по обезличенному тексту, без ключа всё работает на правилах",
                        f"Стенд: {STAND} · код открыт на GitHub"], size=12, color=DARK)
    set_text(ph(s, 0), "АРХИТЕКТУРА И ТЕХНИКА", color=VIOLET)

    # 26 — демо работодателя
    s = S[26]
    set_title(s, "ДЕМО: РАБОТОДАТЕЛЬ")
    picture(s, 14, "11_employer_selection.png", "employer_selection.png")
    set_text(ph(s, 15), "Потребность описывается текстом: специализация, грейд, стек и зарплата распознаются сами", size=10, color=DARK)
    set_text(ph(s, 16), "Подборка по категориям: оценка «подходит», шесть слагаемых и обоснование у каждого кандидата", size=10, color=DARK)
    set_text(ph(s, 17), "Приглашение с зарплатой «от–до»; контакты открываются после согласия кандидата", size=10, color=DARK)

    # 27 — демо кандидата
    s = S[27]
    set_title(s, "ДЕМО: КАНДИДАТ")
    picture(s, 14, "17_test_result.png", "candidate_test.png"); picture(s, 18, "05_candidate_achievements.png", "02_candidate_profile.png")
    set_text(ph(s, 15), "Результат теста: грейд, θ с интервалом, разрез по технологиям; подтверждённые технологии попадают в профиль", size=10, color=DARK)
    set_text(ph(s, 16), "Достижения ФСП: подтверждено источником, уровень, срок действия; ФСП ID связан по OpenID Connect", size=10, color=DARK)

    # 18 — ограничения и план (сетка 6)
    s = S[18]
    set_title(s, "ОГРАНИЧЕНИЯ И ПЛАН")
    grid = [("Синтетика", "метрики — на синтетических мирах с известной истиной и восьми вариантах генератора; на реальных откликах цикл запускается заново"),
            ("Банк и тест", "у четырёх новых специализаций по шесть шаблонов; согласованность 0,50 при 10 заданиях, 0,59 при 15"),
            ("Пилот", "следующий шаг — живые кандидаты и работодатели, сверка грейда с экспертной оценкой, калибровка подбора по откликам"),
            ("ФСП и рынок", "Keycloak и реестр ФСП — три настройки и один адаптер; дальше интеграция с ATS, адаптивная версия, защита от фиктивных вакансий")]
    for (t, d), (ni, ti, di) in zip(grid, [(49, 37, 38), (50, 39, 40), (51, 41, 42), (52, 43, 44)]):
        set_text(ph(s, ni), f"0{grid.index((t, d)) + 1}", bold=True, color=PINK); set_text(ph(s, ti), t, bold=True, color=VIOLET); set_text(ph(s, di), d, size=9.5, color=DARK)

    # 13 — финал: попробовать
    s = S[13]
    set_title(s, "ПОПРОБОВАТЬ")
    set_text(ph(s, 1), [f"Прототип: https://{STAND}/",
                        "Учётные записи для проверки — на экране входа (кандидат с историей ФСП, без истории, работодатель)",
                        f"API: https://{STAND}/api/docs",
                        "Код: github.com/YODA02041994/lct2026-fsp-jobs",
                        f"Документация по восьми разделам ТЗ: https://{STAND}/files/Документация_Карьера_ФСП.pdf",
                        "Спасибо!"], size=13, color=DARK)

    keep_slides(prs, [7, 8, 9, 10, 11, 24, 25, 15, 21, 22, 20, 23, 16, 12, 26, 27, 18, 13])
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    os.makedirs(os.path.join(ROOT, "frontend", "files"), exist_ok=True)
    shutil.copyfile(OUT, os.path.join(ROOT, "frontend", "files", "Карьера_ФСП_презентация.pptx"))
    print("готово:", OUT, "слайдов:", len(prs.slides))


if __name__ == "__main__":
    main()
